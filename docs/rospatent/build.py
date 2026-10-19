# Сборка материалов для заявки в Роспатент (приказ Минэкономразвития № 211)
#  - referat.pdf / referat.docx — реферат, п. 30
#  - deposit.pdf — депонируемые материалы в PDF/A-2b, п. 29: титульный лист
#    и фрагменты исходного текста со сквозной нумерацией страниц и номером
#    страницы полного текста в скобках
# Оформление по п. 14: А4, поля 20/20/20, левое 25 мм, чёрный шрифт 12 пт,
# интервал 1,5, печать на одной стороне, нумерация арабскими цифрами с 1.
#
# Запуск: PYTHONPATH=/tmp/pylibs python3 docs/rospatent/build.py
import datetime
import io
import json
import os
import re

from fpdf import FPDF

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
OUT = os.path.join(ROOT, 'docs', 'rospatent')
FONTS = '/usr/share/fonts/noto'
DATA = json.load(open(os.path.join(OUT, 'data.json'), encoding='utf-8'))

BLANK = '_' * 40

# Сколько страниц полного текста берём из начала и из конца
HEAD_PAGES = 30
TAIL_PAGES = 30

# ---------------------------------------------------------------- исходники

# Свой код программы. Готовые библиотеки (src/components/ui — набор
# shadcn/ui) и сторонние пакеты в депонируемые материалы не входят
GROUPS = [
    ('Программа для Windows (C#)', 'desktop/csharp/PvsPdfApp', ('.cs',)),
    ('Серверная часть (Python)', 'backend', ('.py',)),
    # Интерфейс — последним: конец выдержки приходится на ядро обработки
    # документов, а не на служебные шаблоны писем
    ('Интерфейс и обработка документов (TypeScript)', 'src', ('.ts', '.tsx')),
]

# Порядок внутри групп: сначала то, что лучше всего опознаёт программу
FIRST = [
    'desktop/csharp/PvsPdfApp/Program.cs',
    'desktop/csharp/PvsPdfApp/MainForm.cs',
    'src/main.tsx',
    'src/App.tsx',
    'src/pages/Index.tsx',
    'src/context/DocContext.tsx',
]


def own_files():
    out = []
    for title, base, ext in GROUPS:
        files = []
        for d, _, fs in os.walk(os.path.join(ROOT, base)):
            rel = os.path.relpath(d, ROOT).replace('\\', '/')
            if any(x in rel for x in ('node_modules', '/bin', '/obj', 'src/components/ui')):
                continue
            for f in fs:
                if f.endswith(ext) and f != 'vite-env.d.ts':
                    files.append(f'{rel}/{f}')
        files.sort(key=lambda p: (p not in FIRST, FIRST.index(p) if p in FIRST else 0, p))
        out.append((title, files))
    return out


# Скрываем то, что не нужно публике: адреса серверных функций, ключ
# проверки лицензий, тестовые ключи. На опознание программы это не влияет
MASKS = [
    (re.compile(r'https://functions\.poehali\.dev/[0-9a-f-]+'), 'https://***'),
    (re.compile(r'"[A-Za-z0-9+/=]{40,}"'), '"***"'),
    (re.compile(r"'[A-Za-z0-9+/=]{40,}'"), "'***'"),
]


def clean(line):
    for rx, rep in MASKS:
        line = rx.sub(rep, line)
    return line.replace('\t', '    ').rstrip()


# ---------------------------------------------------------------- оформление

MM = 72 / 25.4


class Doc(FPDF):
    def __init__(self):
        super().__init__(unit='pt', format='A4')
        self.set_margins(25 * MM, 20 * MM, 20 * MM)
        self.set_auto_page_break(False)
        self.add_font('Sans', '', f'{FONTS}/NotoSans-Regular.ttf')
        self.add_font('Sans', 'B', f'{FONTS}/NotoSans-Bold.ttf')
        self.add_font('Mono', '', f'{FONTS}/NotoSansMono-Regular.ttf')
        self.add_font('Mono', 'B', f'{FONTS}/NotoSansMono-Bold.ttf')
        self.set_text_color(0, 0, 0)
        self.set_draw_color(0, 0, 0)
        self.set_creator(DATA['program'])
        self.set_producer(DATA['program'])
        self.set_lang('ru-RU')

    @property
    def width(self):
        return self.w - self.l_margin - self.r_margin


SIZE = 12
LINE = SIZE * 1.5


def holder():
    return DATA.get('holder') or BLANK


def authors():
    return DATA.get('authors') or []


# ---------------------------------------------------------------- реферат

def byte_size():
    total = 0
    for _, files in own_files():
        for f in files:
            total += os.path.getsize(os.path.join(ROOT, f))
    return total


def mb(n):
    return f'{n / 1024 / 1024:.2f}'.replace('.', ',')


def kb(n):
    return f'{round(n / 1024):,}'.replace(',', ' ')


# Должно дословно совпадать с полем «Операционная система» в заявлении
OS_TEXT = ('Windows 10, Windows 11 (64 бит); Astra Linux, РЕД ОС, Альт '
           'и другие ОС — через веб-браузер')


def referat_text(size):
    return (
        f'{DATA["program"]} — редактор документов PDF для компьютеров с Windows '
        'и браузера. Назначение: просмотр, правка, подготовка, печать и '
        'преобразование документов PDF без сторонних программ. Область '
        'применения: делопроизводство и документооборот организаций. '
        'Функциональные возможности: просмотр с поиском и закладками; '
        'надписи, выделение, замечания; добавление, удаление, поворот и '
        'перестановка страниц; объединение и сжатие файлов; преобразование '
        'PDF в Word, Excel, JPG, PNG, текст и HTML и обратно из изображений '
        'и текста; оптическое распознавание русского и английского текста '
        'с созданием PDF с поиском; сканирование со сканеров WIA и TWAIN; '
        'печать; активация лицензии и обновление. Тип ЭВМ: IBM PC-совместимый '
        f'компьютер. ОС: {OS_TEXT}. '
        'Язык: C#, JavaScript (TypeScript), Python. '
        f'Объём программы: {kb(size)} Кбайт.'
    )


def build_referat():
    size = byte_size()
    text = referat_text(size)
    assert len(text) <= 900, f'Реферат длиннее 900 знаков: {len(text)}'

    pdf = Doc()
    pdf.set_title(f'Реферат — {DATA["program"]}')
    pdf.add_page()
    pdf.set_font('Sans', 'B', 14)
    pdf.cell(pdf.width, LINE, 'РЕФЕРАТ', align='C', new_x='LMARGIN', new_y='NEXT')
    pdf.ln(LINE / 2)
    pdf.set_font('Sans', '', SIZE)
    rows = [
        ('Название программы для ЭВМ', DATA['program']),
        ('Правообладатель', holder()),
        ('Автор(ы)', ', '.join(a['name'] for a in authors()) or BLANK),
    ]
    for k, v in rows:
        pdf.set_font('Sans', 'B', SIZE)
        pdf.multi_cell(pdf.width, LINE, k + ':', new_x='LMARGIN', new_y='NEXT')
        pdf.set_font('Sans', '', SIZE)
        pdf.multi_cell(pdf.width, LINE, v, new_x='LMARGIN', new_y='NEXT')
    pdf.ln(LINE / 2)
    pdf.set_font('Sans', '', SIZE)
    pdf.multi_cell(pdf.width, LINE, text, align='J', new_x='LMARGIN', new_y='NEXT')
    pdf.set_y(pdf.h - 20 * MM - LINE)
    pdf.cell(pdf.width, LINE, '1', align='C')
    to_pdfa(bytes(pdf.output()), os.path.join(OUT, 'referat.pdf'), f'Реферат — {DATA["program"]}')

    from docx import Document
    from docx.shared import Mm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin = s.right_margin = s.bottom_margin = Mm(20)
    s.left_margin = Mm(25)
    st = d.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 1.5
    h = d.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run('РЕФЕРАТ')
    r.bold = True
    r.font.size = Pt(14)
    for k, v in rows:
        p = d.add_paragraph()
        p.add_run(k + ': ').bold = True
        p.add_run(v)
    p = d.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.save(os.path.join(OUT, 'referat.docx'))
    return text, size


# ---------------------------------------------------------------- листинг

def wrap(line, cols):
    if not line:
        return ['']
    out = []
    while len(line) > cols:
        out.append(line[:cols])
        line = '  ' + line[cols:]
    out.append(line)
    return out


def listing_pages(pdf):
    # Раскладываем полный исходный текст по страницам заранее: так известен
    # номер каждой страницы в полном тексте, и его можно указать в скобках
    pdf.set_font('Mono', '', SIZE)
    char_w = pdf.get_string_width('M')
    cols = int(pdf.width // char_w)
    body_top = 20 * MM + LINE * 1.6
    body_bottom = pdf.h - 20 * MM - LINE * 1.2
    per_page = int((body_bottom - body_top) // LINE)

    pages = []
    cur = []
    cur_file = None

    def flush():
        nonlocal cur
        if cur:
            pages.append({'file': cur_file, 'lines': cur})
        cur = []

    for title, files in own_files():
        for f in files:
            flush()
            cur_file = f
            cur.append(('head', f'// ===== {f} ====='[:cols]))
            src = open(os.path.join(ROOT, f), encoding='utf-8-sig', errors='replace').read()
            for raw in src.splitlines():
                for part in wrap(clean(raw), cols):
                    if len(cur) >= per_page:
                        flush()
                    cur.append(('code', part))
    flush()
    return pages, cols, body_top


def draw_listing_page(pdf, page, full_no, full_total, body_top):
    pdf.add_page()
    pdf.set_font('Sans', '', 9)
    pdf.set_xy(pdf.l_margin, 20 * MM)
    pdf.cell(pdf.width * 0.6, LINE, f'{DATA["program"]}. Исходный текст программы')
    pdf.cell(pdf.width * 0.4, LINE, page['file'] or '', align='R')
    pdf.line(pdf.l_margin, 20 * MM + LINE, pdf.w - pdf.r_margin, 20 * MM + LINE)

    y = body_top
    for kind, text in page['lines']:
        pdf.set_font('Mono', 'B' if kind == 'head' else '', SIZE)
        pdf.set_xy(pdf.l_margin, y)
        pdf.cell(pdf.width, LINE, text)
        y += LINE

    pdf.set_font('Sans', '', SIZE)
    pdf.set_xy(pdf.l_margin, pdf.h - 20 * MM - LINE)
    pdf.cell(pdf.width, LINE, f'{pdf.page_no()}   ({full_no})', align='C')


def draw_title(pdf, size, full_total, picked):
    pdf.add_page()
    pdf.set_font('Sans', 'B', 16)
    pdf.set_y(60 * MM)
    pdf.multi_cell(pdf.width, 24, 'ДЕПОНИРУЕМЫЕ МАТЕРИАЛЫ,\nИДЕНТИФИЦИРУЮЩИЕ ПРОГРАММУ ДЛЯ ЭВМ', align='C',
                   new_x='LMARGIN', new_y='NEXT')
    pdf.ln(LINE)
    pdf.set_font('Sans', 'B', 20)
    pdf.multi_cell(pdf.width, 30, f'«{DATA["program"]}»', align='C', new_x='LMARGIN', new_y='NEXT')
    pdf.set_font('Sans', '', SIZE)
    pdf.cell(pdf.width, LINE, f'Версия {DATA["version"]}', align='C', new_x='LMARGIN', new_y='NEXT')
    pdf.ln(LINE * 2)

    def row(k, v):
        pdf.set_font('Sans', 'B', SIZE)
        pdf.multi_cell(pdf.width, LINE, k, new_x='LMARGIN', new_y='NEXT')
        pdf.set_font('Sans', '', SIZE)
        pdf.multi_cell(pdf.width, LINE, v, new_x='LMARGIN', new_y='NEXT')
        pdf.ln(LINE / 3)

    row('Правообладатель:', holder())
    names = [a['name'] for a in authors()]
    row('Автор(ы):', '\n'.join(names) if names else f'{BLANK}\n{BLANK}')
    row('Вид материалов:',
        f'Фрагменты исходного текста: страницы 1–{HEAD_PAGES} и '
        f'{full_total - TAIL_PAGES + 1}–{full_total} из {full_total} страниц полного текста. '
        'В скобках внизу каждой страницы указан её номер в полном тексте.')
    row('Языки программирования:', 'C#, JavaScript (TypeScript), Python')
    row('Объём исходного текста:', f'{kb(size)} Кбайт ({size:,} байт)'.replace(',', ' '))

    pdf.set_font('Sans', '', SIZE)
    pdf.set_xy(pdf.l_margin, pdf.h - 20 * MM - LINE * 3)
    city = DATA.get('city') or ''
    pdf.cell(pdf.width, LINE, f'{city + ", " if city else ""}{DATA["year"]}', align='C',
             new_x='LMARGIN', new_y='NEXT')
    pdf.cell(pdf.width, LINE, '1', align='C')


def build_deposit(size):
    pdf = Doc()
    pdf.set_title(f'Депонируемые материалы — {DATA["program"]}')
    pdf.set_subject('Исходный текст программы для ЭВМ')
    pdf.set_author(holder() if DATA.get('holder') else DATA['program'])

    pages, _, body_top = listing_pages(pdf)
    total = len(pages)
    picked = list(range(min(HEAD_PAGES, total)))
    picked += [i for i in range(max(total - TAIL_PAGES, 0), total) if i not in picked]

    draw_title(pdf, size, total, picked)
    for i in picked:
        draw_listing_page(pdf, pages[i], i + 1, total, body_top)

    raw = bytes(pdf.output())
    path = os.path.join(OUT, 'deposit.pdf')
    to_pdfa(raw, path, f'Депонируемые материалы — {DATA["program"]}')
    return total, len(picked) + 1


# ---------------------------------------------------------------- PDF/A

def to_pdfa(raw, path, title):
    import pikepdf
    from PIL import ImageCms

    icc = ImageCms.ImageCmsProfile(ImageCms.createProfile('sRGB')).tobytes()
    with pikepdf.open(io.BytesIO(raw)) as pdf:
        stream = pikepdf.Stream(pdf, icc)
        stream['/N'] = 3
        pdf.Root.OutputIntents = pikepdf.Array([
            pikepdf.Dictionary(
                Type=pikepdf.Name.OutputIntent,
                S=pikepdf.Name.GTS_PDFA1,
                OutputConditionIdentifier=pikepdf.String('sRGB IEC61966-2.1'),
                Info=pikepdf.String('sRGB IEC61966-2.1'),
                DestOutputProfile=stream,
            )
        ])
        pdf.Root.MarkInfo = pikepdf.Dictionary(Marked=True)
        now = datetime.datetime.now(datetime.timezone.utc).astimezone()
        with pdf.open_metadata(set_pikepdf_as_editor=False) as meta:
            meta['pdfaid:part'] = '2'
            meta['pdfaid:conformance'] = 'B'
            meta['dc:title'] = title
            meta['dc:creator'] = [DATA.get('holder') or DATA['program']]
            meta['xmp:CreatorTool'] = DATA['program']
            meta['pdf:Producer'] = DATA['program']
            meta['xmp:CreateDate'] = now.isoformat(timespec='seconds')
            meta['xmp:ModifyDate'] = now.isoformat(timespec='seconds')
        # Сведения в словаре документа должны совпадать с XMP
        stamp = now.strftime("D:%Y%m%d%H%M%S") + now.strftime('%z')[:3] + "'" + now.strftime('%z')[3:] + "'"
        pdf.docinfo['/Title'] = title
        pdf.docinfo['/Author'] = DATA.get('holder') or DATA['program']
        pdf.docinfo['/Creator'] = DATA['program']
        pdf.docinfo['/Producer'] = DATA['program']
        pdf.docinfo['/CreationDate'] = stamp
        pdf.docinfo['/ModDate'] = stamp
        for k in ('/Subject', '/Keywords'):
            if k in pdf.docinfo:
                del pdf.docinfo[k]
        pdf.save(path, fix_metadata_version=False, deterministic_id=True,
                 object_stream_mode=pikepdf.ObjectStreamMode.preserve)


if __name__ == '__main__':
    text, size = build_referat()
    total, sent = build_deposit(size)
    print('Реферат, знаков:', len(text))
    print('Объём исходного текста, байт:', size)
    print('Полный текст, страниц:', total, '| в материалах страниц:', sent)